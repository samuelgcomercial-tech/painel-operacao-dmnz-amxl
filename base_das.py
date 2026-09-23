# -*- coding: utf-8 -*-
"""
base_das.py (versao web)

Mesma logica do trecho "DAs DMNZ" do processa_csv_scc.py do robo desktop
(funcao revisa_cadastro_das) - so devolve os dados pra tela decidir o que
perguntar (via botoes/pills), em vez de usar input()/print() de terminal.
A base de DAs em si mora no GitHub (ve github_store.py) - esse arquivo so
cuida do "CSV em texto" pra dentro/fora, sem saber de onde ele veio.

Le o CSV bruto exportado do SCC (Pesquisar -> Export to CSV), acha quais
motoristas ("Last Scan By") rodaram hoje e ainda nao estao cadastrados
como DMNZ ou nao, e monta a base atualizada depois que o usuario escolhe.

Ainda NAO faz (fica pra depois, de proposito):
  - a excecao "esse DA nao rodou pela DMNZ hoje" (so por hoje / permanente)
  - a coluna "State Finalizador" (justificativa por TBR)
  - gerar o CSV final com as colunas adicionadas
"""

import csv
import io

# Nomes de sistema/robo que aparecem no "Last Scan By" no lugar de um DA de
# verdade (ex: quando a leitura foi feita por automacao, nao por uma
# pessoa) - mesma lista do robo desktop.
NOMES_SISTEMA = {"LASTMILEROUTEPLANNER"}


def normaliza_nome(nome):
    """Deixa o nome tolerante a espaco duplo, espaco nas pontas e diferenca
    de maiuscula/minuscula - os nomes que chegam do SCC vem bem
    inconsistentes nesse quesito. Mesma funcao do robo desktop."""
    return " ".join(nome.split()).upper()


def eh_da_de_verdade(last_scan_by):
    """Mesma regra do robo desktop (login de suporte com @, ou nome de
    sistema, nao conta como DA de verdade) + dois casos que o Samuel
    apontou, os dois com a mesma explicacao de fundo: o SCC registra um
    valor de SISTEMA no Last Scan By enquanto o pacote ainda nao tem um
    motorista de verdade confirmado - o nome real só aparece na leitura
    SEGUINTE (a de entrega, ou a de falha na entrega). Até lá, esse TBR
    fica sem DA de verdade vinculado (mesmo tratamento que já existia
    pro login de suporte), em vez de arriscar cravar um motorista
    errado:

      - login comecando com "P2P" (ex:
        "P2PTransportRequestAssignmentService") - rota dividida, outro
        motorista "puxou" o pacote pra si no meio do caminho.
      - "None" (igualdade EXATA, nao "contém" - pra nao arriscar
        excluir por engano um motorista de verdade cujo nome tenha essa
        sequencia de letras) - pacote de NA (nao foi manifestado/nao foi
        recebido nem atribuido a rota pela equipe da noite, que nem
        consegue nesse caso) cuja rota foi atribuida direto pela equipe
        da Amazon."""
    if not last_scan_by:
        return False
    if "@" in last_scan_by:
        return False
    if last_scan_by.strip().upper().startswith("P2P"):
        return False
    if last_scan_by.strip().lower() == "none":
        return False
    if normaliza_nome(last_scan_by) in NOMES_SISTEMA:
        return False
    return True


def _detecta_delimitador(texto_amostra):
    """Excel com configuracao regional BR costuma salvar CSV com ';' em vez
    de ',' - deixa a leitura tolerante aos dois. Mesma ideia do robo
    desktop, so que a partir de um texto ja em memoria em vez de reabrir
    o arquivo do disco."""
    return ";" if texto_amostra.count(";") > texto_amostra.count(",") else ","


def _eh_arquivo_excel(bytes_arquivo):
    """Detecta se os bytes sao um arquivo Excel (.xlsx/.xls novo, base
    ZIP) pela ASSINATURA BINARIA, nao pelo nome do arquivo - o usuario
    pode chamar de '.csv' um arquivo que na verdade e um Excel (ou o
    contrario), entao confiar so na extensao arrisca ler errado."""
    return bytes_arquivo[:4] == b"PK\x03\x04"


def _le_scc_de_excel(bytes_arquivo):
    """Le um CSV do SCC que foi salvo como .xlsx em vez de .csv - visto
    de verdade em 23/09/2026 (time noturno abriu o CSV exportado do SCC
    no Excel com separador de lista regional diferente de virgula, o
    Excel nao separou as colunas, e salvaram assim mesmo). Cobre os
    dois formatos possiveis:

      - tabela normal (cabecalho na linha 1, uma coluna por campo) -
        caso alguem exporte/salve certinho no futuro.
      - "CSV disfarcado de xlsx" (o caso real visto): cada linha
        inteira do CSV original virou UM texto so dentro da coluna A,
        aspas e virgulas do CSV original preservadas dentro do texto -
        remonta o texto original e reusa o parser de CSV de sempre.

    Nos dois casos devolve a mesma lista de dicts que le_csv_scc."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(bytes_arquivo), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    linhas = list(ws.iter_rows(values_only=True))
    if not linhas:
        return []

    cabecalho = linhas[0]
    colunas_preenchidas = sum(1 for c in cabecalho if c not in (None, ""))
    if colunas_preenchidas > 1:
        nomes = [str(c or "").strip() for c in cabecalho]
        return [
            {nomes[i]: row[i] for i in range(len(nomes)) if i < len(row)}
            for row in linhas[1:]
        ]

    texto_csv = "\n".join(str(row[0]) for row in linhas if row and row[0] is not None)
    delimitador = _detecta_delimitador(texto_csv[:2000])
    reader = csv.DictReader(io.StringIO(texto_csv), delimiter=delimitador)
    return list(reader)


def _decodifica_texto_csv(bytes_arquivo):
    """Decodifica os bytes do CSV pra texto, tolerando encoding diferente
    de UTF-8 - visto de verdade em 24/09/2026 (Samuel exportando pelo
    SCC no celular): o arquivo baixado no celular vem com nome '.xlsx'
    mas o CONTEUDO e texto CSV puro (nao passa no teste de assinatura
    binaria do Excel, ver _eh_arquivo_excel) - so que em Windows-1252
    (cp1252), nao UTF-8, porque assim que o SCC/navegador do celular
    exporta. UTF-8 e cp1252 tem baixa compatibilidade cruzada pra
    acentos: um 'ã' sozinho em cp1252 e o byte 0xE3, que o decodificador
    UTF-8 rejeita direto ('invalid continuation byte') por nao formar
    uma sequencia UTF-8 valida - foi exatamente o erro que apareceu.

    Tenta UTF-8 primeiro (com BOM opcional - continua sendo o formato
    mais comum vindo do desktop); se falhar, tenta cp1252 (cobre
    acentuacao de Windows/BR, nunca falha em decodificar um arquivo
    de verdade porque cp1252 mapeia praticamente todo byte); como
    ultimo recurso (nunca deveria chegar aqui) usa latin-1 substituindo
    o que nao der pra decodificar, pra nunca travar o robo por causa
    de encoding."""
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return bytes_arquivo.decode(encoding)
        except UnicodeDecodeError:
            continue
    return bytes_arquivo.decode("latin-1", errors="replace")


def le_csv_scc(arquivo):
    """Le o CSV exportado do SCC a partir de um objeto de upload do
    Streamlit. Devolve lista de dicts (uma por linha), igual csv.DictReader
    do robo desktop.

    Tambem aceita esse mesmo export salvo como .xlsx (ver
    _le_scc_de_excel) - detectado pela assinatura binaria do arquivo,
    nao pela extensao, entao funciona mesmo se o nome do arquivo
    continuar dizendo '.csv' (ou, no caso oposto, '.xlsx' sem ser Excel
    de verdade - ver _decodifica_texto_csv)."""
    bytes_arquivo = arquivo.getvalue()
    if _eh_arquivo_excel(bytes_arquivo):
        return _le_scc_de_excel(bytes_arquivo)
    texto = _decodifica_texto_csv(bytes_arquivo)
    delimitador = _detecta_delimitador(texto[:2000])
    reader = csv.DictReader(io.StringIO(texto), delimiter=delimitador)
    return list(reader)


def extrai_das_do_dia(linhas_csv):
    """Lista (sem repetir, em ordem alfabetica) os nomes de DA de verdade
    que aparecem no Last Scan By das linhas de hoje. Mesma funcao do robo
    desktop."""
    nomes = set()
    for l in linhas_csv:
        nome = (l.get("Last Scan By") or "").strip()
        if eh_da_de_verdade(nome):
            nomes.add(nome)
    return sorted(nomes)


def parseia_texto_base(texto):
    """Le a base de DAs a partir de um texto CSV já em memória (Nome do
    DA,DMNZ) - usado tanto pro que vem do GitHub quanto (se um dia
    precisar de novo) de um upload manual. Se o texto vier vazio/None
    (primeira vez, arquivo ainda não existe), devolve tudo vazio - não
    trava o app, só significa que todo mundo vai aparecer como "ainda
    não cadastrado".

    Devolve (base_dict, linhas_bruta):
      - base_dict: nome normalizado -> classificacao (ex.: "DMNZ")
      - linhas_bruta: lista de (nome_original, classificacao), preserva a
        grafia original (acentos/maiusculas) pra nao perder formatacao ao
        salvar de novo.
    """
    if not texto:
        return {}, []

    delimitador = _detecta_delimitador(texto[:2000])
    reader = csv.DictReader(io.StringIO(texto), delimiter=delimitador)
    fieldnames_norm = [(c or "").strip().lower() for c in (reader.fieldnames or [])]
    if "nome do da" not in fieldnames_norm or "dmnz" not in fieldnames_norm:
        raise ValueError(
            "A base de DAs no repositório não tem as colunas esperadas "
            f"'Nome do DA' e 'DMNZ' (encontrei: {reader.fieldnames})."
        )
    col_nome = reader.fieldnames[fieldnames_norm.index("nome do da")]
    col_dmnz = reader.fieldnames[fieldnames_norm.index("dmnz")]

    base_dict = {}
    linhas_bruta = []
    for row in reader:
        nome_bruto = row.get(col_nome)
        if not nome_bruto:
            continue
        # protecao: nome contaminado com virgula colada na mesma celula
        # (erro comum de digitacao no Excel) - usa só a parte antes dela.
        if "," in nome_bruto:
            nome_bruto = nome_bruto.split(",")[0]
        classificacao = (row.get(col_dmnz) or "").strip()
        base_dict[normaliza_nome(nome_bruto)] = classificacao
        linhas_bruta.append((nome_bruto.strip(), classificacao))

    return base_dict, linhas_bruta


def gera_csv_base_das(linhas_bruta):
    """Monta o texto do CSV da base de DAs (mesmo formato do robo desktop:
    'Nome do DA,DMNZ'), pronto pra oferecer como download."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["Nome do DA", "DMNZ"])
    writer.writerows(linhas_bruta)
    return buffer.getvalue()
