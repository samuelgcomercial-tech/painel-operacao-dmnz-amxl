# -*- coding: utf-8 -*-
"""
exporta_fechamento.py (versao web)

Gera o Excel final do Fechamento (Etapa 3) pra download - pedido do
Samuel em 13/09/2026. Planejado em 3 abas, construidas uma de cada vez:

  1. "Analise do dia" (leva 1, pronta): uma linha por TBR, com TODAS
     as colunas do CSV original do SCC + duas colunas calculadas no
     final (State DA's e State Finalizador). E a "base" pronta pra ele
     montar tabela dinamica de verdade em cima dela no Excel (criar a
     tabela dinamica em si via codigo exigiria mexer direto no XML
     interno do arquivo - fragil e dificil de manter; a base "achatada"
     e o que da trabalho de verdade pra montar a mao, e isso o robo ja
     automatiza).
  2. "Apresentacao" (leva 2, pronta - pedido do Samuel em 14/09/2026):
     "replica" estatica dos 5 paineis que ja saem no dashboard HTML
     (calculo_fechamento.calcula_tudo / dashboard_fechamento.py),
     simulando o visual de tabela dinamica (grupos colapsaveis via
     outline do Excel) ja que uma Tabela Dinamica de verdade nao da
     pra gerar por codigo sem editar o XML interno na mao.
  3. "Historico" (ultima leva, ainda NAO feita): export do
     historico_tbr.py (base persistida no GitHub) pra auditoria -
     desde quando cada TBR problematico foi visto e quando foi
     verificado pela ultima vez.
"""

import datetime as dt
import io

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from state_finalizador import monta_coluna_das_dmnz

# Mesma ordem/nomes de coluna do CSV cru exportado pelo SCC - conferido
# com o arquivo real (SearchResults_4.csv) em 13/09/2026.
COLUNAS_CSV_ORIGINAL = [
    "Tracking ID", "Last Updated Time", "Source", "State", "Destination",
    "Reason", "Station", "Operation", "Route Code", "Last Scan By",
    "Address Type", "Ship Option", "Ship Method", "Delivery Destination Type",
]

COL_STATE_DAS = "State DA's"
COL_STATE_FINALIZADOR = "State Finalizador"

# Coluna de data/hora do CSV cru - pedido do Samuel em 13/09/2026: mostrar
# so dd/mm/aaaa, sem hh:mm:ss. O valor continua guardando a hora por
# baixo dos panos (so o FORMATO de exibicao muda) - assim, se um dia
# precisar ordenar/filtrar pela hora exata, o dado nao foi perdido, so
# nao aparece na tela por padrao.
COL_DATA_HORA = "Last Updated Time"
FORMATO_DATA_HORA_ORIGINAL = "%Y-%m-%d %H:%M:%S"


def _parseia_data_hora(texto):
    """Converte o texto do CSV ('2026-09-13 07:39:52') pra um datetime de
    verdade, pra o Excel tratar como data (e nao como texto solto). Se
    vier vazio ou num formato inesperado, devolve o texto original sem
    quebrar - melhor mostrar o dado cru do que sumir com a linha."""
    texto = (texto or "").strip()
    if not texto:
        return ""
    try:
        return dt.datetime.strptime(texto, FORMATO_DATA_HORA_ORIGINAL)
    except ValueError:
        return texto


def monta_linhas_analise_do_dia(linhas_csv, base_das, motorista_real_da_rota, base_tbr):
    """Uma linha por TBR: todas as colunas originais do CSV do SCC + State
    DA's (DMNZ ou vazio - mesma fonte que ja alimenta a tela,
    monta_coluna_das_dmnz) + State Finalizador (mesma classificacao que
    ja sai no dashboard: 'Entregue' automatico pra quem foi Delivered,
    ou o que estiver salvo no historico pra quem nao foi - mesma regra
    de calculo_fechamento.monta_dados_do_dia, so que aqui devolvendo
    linha completa em vez de so os 4 campos resumidos)."""
    dmnz_por_tbr = monta_coluna_das_dmnz(linhas_csv, base_das, motorista_real_da_rota)

    linhas = []
    for l in linhas_csv:
        tbr = l["Tracking ID"]
        state_scc = (l.get("State") or "").strip()
        if state_scc.lower() == "delivered":
            finalizador = "Entregue"
        else:
            entrada = base_tbr.get(tbr)
            finalizador = entrada["classificacao"] if entrada else ""

        linha = {col: l.get(col, "") for col in COLUNAS_CSV_ORIGINAL}
        linha[COL_DATA_HORA] = _parseia_data_hora(linha[COL_DATA_HORA])
        linha[COL_STATE_DAS] = dmnz_por_tbr.get(tbr, "")
        linha[COL_STATE_FINALIZADOR] = finalizador
        linhas.append(linha)
    return linhas


LIMITE_TBRS_NA_FRASE = 10


def _lista_tbrs_ate_limite(tbrs):
    """Formata uma lista de TBRs pra entrar dentro da frase do resumo -
    mesmo padrao ja usado em calculo_fechamento.monta_observacoes pro
    'PCT NA' (lista os primeiros 10, com '...' se sobrar mais que
    isso). So entra na frase quando a lista nao esta vazia."""
    if not tbrs:
        return ""
    texto = ", ".join(tbrs[:LIMITE_TBRS_NA_FRASE])
    if len(tbrs) > LIMITE_TBRS_NA_FRASE:
        texto += ", ..."
    return f" ({texto})"


def monta_linhas_resumo(linhas, resultado_calculo, tbrs_ja_conhecidos, nome_usuario, data_criacao_str):
    """Monta as frases do resumo no topo da aba, mesmo padrao do robo
    SSW 081 (celula mesclada, uma frase por linha). Pedido do Samuel em
    13/09/2026 (13/09 tambem pediu pra listar os TBRs revalidados
    automatico na propria frase, pra dar pra auditar sem abrir mais
    nada).

    Contas:
      - entregues: contagem_scc['Delivered'], ja calculado pelo
        calculo_fechamento.calcula_tudo.
      - em_rota: quem ficou com State Finalizador comecando com 'EM
        ROTA' (classificacao automatica de In Transit/Failed com DA de
        verdade vinculado - ver state_finalizador.py) - ainda nao
        confirmado como entregue nem como insucesso.
      - revalidados sem alteracao: tbrs_ja_conhecidos, que PRECISA vir
        de fora (Etapa 2, historico_tbr.decide_quem_precisa_reabrir) -
        nao da pra saber só pelo resultado final quem reaproveitou a
        classificacao de ontem e quem foi digitado hoje, os dois
        terminam com o mesmo texto salvo.
      - validados por pessoa hoje: o resto de quem nao foi entregue
        nem ficou em EM ROTA automatico nem foi reaproveitado do
        historico - ou seja, quem passou por revisao manual de
        verdade nesta rodada.
      - insucesso sem parceiro: contagem_insucesso['sem_detalhe'], ja
        calculado pelo calculo_fechamento.calcula_tudo (usuario digitou
        so 'Insucesso', sem citar DMNZ/nome de parceiro cadastrado)."""
    total = len(linhas)
    entregues = resultado_calculo["contagem_scc"].get("Delivered", 0)
    em_rota = sum(
        1 for l in linhas if l.get(COL_STATE_FINALIZADOR, "").upper().startswith("EM ROTA")
    )
    n_ja_conhecidos = len(tbrs_ja_conhecidos)
    nao_entregues = total - entregues
    validados_hoje = max(nao_entregues - em_rota - n_ja_conhecidos, 0)
    insucesso_sem_parceiro = resultado_calculo["contagem_insucesso"].get("sem_detalhe", 0)

    return [
        f"{entregues} pacotes entregues",
        f"{n_ja_conhecidos} pacotes revalidados com state sem alteração"
        + _lista_tbrs_ate_limite(tbrs_ja_conhecidos),
        f"{validados_hoje} pacotes validados por {nome_usuario} dia {data_criacao_str}",
        f"{em_rota} pacotes ainda em rota",
        f"{insucesso_sem_parceiro} pacotes de insucesso sem parceiro atrelado",
    ]


def escreve_resumo_topo_mesclado(ws, n_colunas, linhas_resumo):
    """Escreve o resumo numa unica celula mesclada cobrindo da coluna A
    ate a ultima coluna da tabela, uma linha de planilha por frase
    (quebra de linha dentro da propria celula). Mesmo padrao ja usado
    no robo SSW 081 (escreve_resumo_topo_mesclado de
    processa_sswweb.py). Devolve quantas linhas de planilha ocupou, pra
    quem chama saber em que linha o cabecalho da tabela deve comecar."""
    linhas_resumo = [l for l in linhas_resumo if l]
    if not linhas_resumo:
        return 0
    n_linhas = len(linhas_resumo)
    if n_linhas > 1:
        ws.merge_cells(start_row=1, start_column=1, end_row=n_linhas, end_column=max(n_colunas, 1))
    cell = ws.cell(row=1, column=1, value="\n".join(linhas_resumo))
    cell.font = Font(bold=True)
    cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
    for i in range(1, n_linhas + 1):
        ws.row_dimensions[i].height = 15.5
    return n_linhas


def escreve_aba_analise_do_dia(wb, linhas, linhas_resumo=None):
    """Escreve a aba 'Analise do dia', formatada como Tabela de verdade
    do Excel (nao so texto em negrito): resumo mesclado no topo (se
    'linhas_resumo' vier preenchido), cabecalho da tabela com
    filtro/ordenacao automatico, linhas zebradas pelo estilo da tabela,
    congelado (linha do cabecalho fixa ao rolar), sem linha de grade,
    largura de coluna ajustada ao conteudo. Pedido do Samuel em
    13/09/2026. Isso tambem deixa a base pronta pra virar Tabela
    Dinamica de verdade no Excel (Inserir > Tabela Dinamica), ja que
    uma tabela nomeada e a fonte de dados ideal pra isso - cresce
    sozinha se amanha vier mais linha."""
    colunas = COLUNAS_CSV_ORIGINAL + [COL_STATE_DAS, COL_STATE_FINALIZADOR]
    ws = wb.create_sheet("Analise do dia")
    ws.sheet_view.showGridLines = False

    n_linhas_resumo = escreve_resumo_topo_mesclado(ws, len(colunas), linhas_resumo or [])
    linha_cabecalho = n_linhas_resumo + 2 if n_linhas_resumo else 1

    for col_idx, nome_col in enumerate(colunas, start=1):
        cel = ws.cell(row=linha_cabecalho, column=col_idx, value=nome_col)
        cel.font = Font(bold=True)
        cel.alignment = Alignment(horizontal="center")

    for offset, linha in enumerate(linhas):
        row_idx = linha_cabecalho + 1 + offset
        for col_idx, nome_col in enumerate(colunas, start=1):
            valor = linha.get(nome_col, "")
            cel = ws.cell(row=row_idx, column=col_idx, value=valor)
            if nome_col == COL_DATA_HORA and isinstance(valor, dt.datetime):
                cel.number_format = "DD/MM/YYYY"

    ws.freeze_panes = f"A{linha_cabecalho + 1}"

    for col_idx, nome_col in enumerate(colunas, start=1):
        if nome_col == COL_DATA_HORA:
            # largura fixa pro tamanho de "31/12/2026" exibido, nao do
            # texto original ("2026-09-13 07:39:52", mais longo) que nao
            # aparece mais na tela.
            largura = max(len(nome_col), 12)
        else:
            largura = max(
                [len(str(nome_col))] + [len(str(l.get(nome_col, ""))) for l in linhas]
            )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(largura + 2, 40)

    # Tabela de verdade (ListObject) so faz sentido com pelo menos 1
    # linha de dado - com a base vazia, fica so o cabecalho formatado
    # (nao quebra, so nao vira Tabela).
    if linhas:
        ultima_coluna = get_column_letter(len(colunas))
        ultima_linha = linha_cabecalho + len(linhas)
        tabela = Table(
            displayName="Tab_AnaliseDoDia",
            ref=f"A{linha_cabecalho}:{ultima_coluna}{ultima_linha}",
        )
        tabela.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium9", showRowStripes=True,
            showFirstColumn=False, showLastColumn=False, showColumnStripes=False,
        )
        ws.add_table(tabela)

    return ws


# --- Aba "Apresentacao" -----------------------------------------------
#
# "Replica" estatica dos 5 paineis que ja saem no dashboard HTML
# (dashboard_fechamento.gera_html_fechamento) - pedido do Samuel em
# 14/09/2026. Nao da pra gerar uma Tabela Dinamica de verdade por
# codigo sem mexer direto no XML interno (pivotCache/pivotTable -
# fragil e dificil de manter, ver comentario no topo do arquivo), entao
# essa aba SIMULA o resultado: mesmos numeros, mesmo agrupamento por
# categoria (com os grupos aninhados do painel 3 colapsaveis via
# outline do Excel, pra chegar mais perto da sensacao de tabela
# dinamica). E uma FOTOGRAFIA deste fechamento - nao recalcula sozinha
# se a aba "Analise do dia" for editada depois.

COR_LARANJA_HEX = "FFFD4701"  # mesma cor de cabecalho de painel do dashboard_fechamento.py
COR_NAVY_HEX = "FF000D2D"      # mesma cor do cabecalho geral do dashboard_fechamento.py
COR_CATEGORIA_HEX = "FFE8E8E8"  # cinza claro so pra destacar a linha-categoria do painel 3

BORDA_TOPO = Border(top=Side(style="thin"))


# Layout em 3 blocos de colunas lado a lado - mesmo esquema do grid de 3
# colunas do dashboard HTML (.fech-grid: col1 | col2 | col3), pedido do
# Samuel em 14/09/2026: o painel 3 (nested, o que mais parece tabela
# dinamica) fica CENTRAL e mais largo, com os paineis menores nas
# colunas estreitas dos dois lados - painel 1+2 na esquerda, painel
# 4+5 na direita. Cada bloco e "label" (col N) + "valor" (col N+1);
# uma coluna em branco separa um bloco do outro.
COL_ESQUERDA = 1  # A/B - paineis 1 e 2
COL_CENTRO = 4    # D/E - painel 3 (nested)
COL_DIREITA = 7   # G/H - paineis 4 e 5


def _preenche_cabecalho_painel(ws, row, titulo, col_ini=COL_ESQUERDA):
    """Titulo de painel (ex: '1. STATE SCC'), mesclado nas 2 colunas do
    bloco (col_ini e col_ini+1), fundo laranja e fonte branca - mesmo
    visual do cabecalho de painel do dashboard HTML."""
    ws.merge_cells(start_row=row, start_column=col_ini, end_row=row, end_column=col_ini + 1)
    cel = ws.cell(row=row, column=col_ini, value=titulo)
    cel.font = Font(bold=True, color="FFFFFFFF")
    cel.fill = PatternFill(fill_type="solid", fgColor=COR_LARANJA_HEX)
    cel.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 18
    return row + 1


def _escreve_linha_painel(
    ws, row, label, valor, col_ini=COL_ESQUERDA, indent=0, bold=False,
    fill=None, borda_topo=False, outline_level=0,
):
    """Uma linha 'label | valor' dentro de um painel, no bloco de colunas
    col_ini/col_ini+1. indent>0 simula a sub-linha indentada do
    dashboard (.linha.sub); outline_level>0 marca a linha como detalhe
    agrupavel (colapsavel no Excel, grupo Dados > Agrupar) - usado nas
    sub-linhas do painel 3."""
    cel_label = ws.cell(row=row, column=col_ini, value=label)
    cel_label.alignment = Alignment(horizontal="left", indent=indent)
    cel_valor = ws.cell(row=row, column=col_ini + 1, value=valor)
    cel_valor.alignment = Alignment(horizontal="right")
    if bold:
        cel_label.font = Font(bold=True)
        cel_valor.font = Font(bold=True)
    if fill:
        cel_label.fill = fill
        cel_valor.fill = fill
    if borda_topo:
        cel_label.border = BORDA_TOPO
        cel_valor.border = BORDA_TOPO
    if outline_level:
        ws.row_dimensions[row].outlineLevel = outline_level
    return row + 1


def _escreve_painel_state_scc(ws, row, contagem_scc, total_geral, col_ini=COL_ESQUERDA):
    row = _preenche_cabecalho_painel(ws, row, "1. STATE SCC", col_ini)
    for cat, qtd in contagem_scc.most_common():
        row = _escreve_linha_painel(ws, row, cat, qtd, col_ini)
    row = _escreve_linha_painel(ws, row, "Total Geral", total_geral, col_ini, bold=True, borda_topo=True)
    return row + 1  # linha em branco de espacamento ate o proximo painel


def _escreve_painel_entregas_dmnz(ws, row, n_dmnz, col_ini=COL_ESQUERDA):
    row = _preenche_cabecalho_painel(ws, row, "2. STATE ENTREGAS DMNZ", col_ini)
    row = _escreve_linha_painel(ws, row, "DMNZ", n_dmnz, col_ini)
    return row + 1


FILL_CATEGORIA = PatternFill(fill_type="solid", fgColor=COR_CATEGORIA_HEX)


def _escreve_painel_nested(ws, row, nested, total_geral, col_ini=COL_CENTRO):
    row = _preenche_cabecalho_painel(ws, row, "3. STATE SCC & ANÁLISE DMNZ", col_ini)
    for cat, subitens in nested.items():
        row = _escreve_linha_painel(ws, row, cat, "", col_ini, bold=True, fill=FILL_CATEGORIA)
        for fin, qtd in subitens.items():
            row = _escreve_linha_painel(
                ws, row, fin or "(sem State Finalizador)", qtd, col_ini, indent=1, outline_level=1
            )
    row = _escreve_linha_painel(ws, row, "Total Geral", total_geral, col_ini, bold=True, borda_topo=True)
    return row + 1


def _escreve_painel_insucesso(ws, row, contagem_insucesso, col_ini=COL_DIREITA):
    row = _preenche_cabecalho_painel(ws, row, "4. ANÁLISE INSUCESSO", col_ini)
    # mesma ordenacao do dashboard: "sem_detalhe" por ultimo de proposito
    # (e o caso que precisa de atencao, nao compete por ordem com os
    # nomes de parceiro de verdade).
    itens = sorted(
        contagem_insucesso.items(), key=lambda kv: (kv[0] == "sem_detalhe", -kv[1])
    )
    if itens:
        for nome, qtd in itens:
            label = "Insucesso (sem detalhe)" if nome == "sem_detalhe" else f"Insucesso {nome}"
            row = _escreve_linha_painel(ws, row, label, qtd, col_ini)
    else:
        row = _escreve_linha_painel(ws, row, "Sem registros", "-", col_ini)
    return row + 1


def _escreve_painel_mnr(ws, row, total_mnr, col_ini=COL_DIREITA):
    row = _preenche_cabecalho_painel(ws, row, "5. MNR's", col_ini)
    if total_mnr:
        row = _escreve_linha_painel(ws, row, "MNR", total_mnr, col_ini)
    else:
        row = _escreve_linha_painel(ws, row, "Sem registros", "-", col_ini)
    return row + 1


def escreve_aba_apresentacao(wb, resultado_calculo, node, data_fechamento):
    """Segunda aba: os mesmos 5 paineis, agora em 3 blocos de coluna lado
    a lado (esquerda | centro | direita), igual o grid de 3 colunas do
    dashboard HTML - painel 3 (o "tabela dinamica" simulado) central e
    mais largo, paineis 1+2 na coluna estreita da esquerda, paineis
    4+5 na coluna estreita da direita. Cada bloco cresce pra baixo de
    forma independente (alturas diferentes normais - o painel 3 quase
    sempre e o mais alto). 'resultado_calculo' e o mesmo dict devolvido
    por calculo_fechamento.calcula_tudo (o mesmo que ja alimenta o
    dashboard HTML - nunca recalcula nada diferente aqui)."""
    ws = wb.create_sheet("Apresentação")
    ws.sheet_view.showGridLines = False
    # Grupo do painel 3 colapsa/expande pelo BOTAO ACIMA do grupo (onde
    # fica a linha-categoria), nao abaixo - contrario do padrao do
    # Excel (resumo embaixo do detalhe), porque aqui o cabecalho da
    # categoria vem ANTES das sub-linhas, nao depois.
    ws.sheet_properties.outlinePr.summaryBelow = False

    row_titulo = 1
    titulo = f"Fechamento {node} – {data_fechamento.strftime('%d/%m/%Y')}"
    ws.merge_cells(start_row=row_titulo, start_column=1, end_row=row_titulo, end_column=8)
    cel = ws.cell(row=row_titulo, column=1, value=titulo)
    cel.font = Font(bold=True, color="FFFFFFFF")
    cel.fill = PatternFill(fill_type="solid", fgColor=COR_NAVY_HEX)
    cel.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row_titulo].height = 20

    row_inicio = row_titulo + 2
    r = resultado_calculo

    row = row_inicio
    row = _escreve_painel_state_scc(ws, row, r["contagem_scc"], r["total_geral"], COL_ESQUERDA)
    row = _escreve_painel_entregas_dmnz(ws, row, r["n_dmnz"], COL_ESQUERDA)

    row = row_inicio
    row = _escreve_painel_nested(ws, row, r["nested"], r["total_geral"], COL_CENTRO)

    row = row_inicio
    row = _escreve_painel_insucesso(ws, row, r["contagem_insucesso"], COL_DIREITA)
    row = _escreve_painel_mnr(ws, row, r["total_mnr"], COL_DIREITA)

    larguras = {
        "A": 24, "B": 10,   # esquerda - estreita
        "C": 3,             # espacador
        "D": 34, "E": 12,   # centro - mais larga (painel "tabela dinamica")
        "F": 3,             # espacador
        "G": 22, "H": 10,   # direita - estreita
    }
    for letra, largura in larguras.items():
        ws.column_dimensions[letra].width = largura
    return ws


def gera_workbook_fechamento(
    linhas_csv, base_das, motorista_real_da_rota, base_tbr,
    resultado_calculo, tbrs_ja_conhecidos, nome_usuario, data_criacao_str,
    node, data_fechamento,
):
    """Ponto de entrada. Duas abas ja prontas ('Analise do dia' e
    'Apresentacao') - a aba 'Historico' entra na proxima leva, depois
    de validar essas duas.

    resultado_calculo: dict devolvido por calculo_fechamento.calcula_tudo
    (usa contagem_scc e contagem_insucesso pro resumo da aba 1, e todos
    os campos pra aba 2).
    tbrs_ja_conhecidos: lista de TBRs (strings) que reaproveitaram a
    classificacao de ontem sem revisao nova hoje (vem da Etapa 2,
    historico_tbr.decide_quem_precisa_reabrir - a lista 'ja_conhecidos'
    de la, so os TBRs; nao da pra calcular aqui dentro).
    nome_usuario: st.session_state.nome_usuario (quem esta logado no
    app nessa sessao).
    data_criacao_str: data de hoje ja formatada 'dd/mm/aaaa' (dia em
    que o Excel foi gerado - usada na frase do resumo da aba 1).
    node: NODE_ATUAL do app.py (ex: 'LRN9') - so pro titulo da aba 2.
    data_fechamento: data do ARQUIVO de rotas (st.session_state.
    data_arquivo_rotas), mesma usada no titulo do dashboard HTML e no
    nome do arquivo - so pro titulo da aba 2.

    Devolve bytes prontos pro st.download_button do app.py."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # tira a aba padrao "Sheet" vazia

    linhas_base = monta_linhas_analise_do_dia(linhas_csv, base_das, motorista_real_da_rota, base_tbr)
    linhas_resumo = monta_linhas_resumo(
        linhas_base, resultado_calculo, tbrs_ja_conhecidos, nome_usuario, data_criacao_str
    )
    escreve_aba_analise_do_dia(wb, linhas_base, linhas_resumo)
    escreve_aba_apresentacao(wb, resultado_calculo, node, data_fechamento)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

